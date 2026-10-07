"""
Customer purchase data categorization (sample project, fictional data)

Takes a messy purchase export (inconsistent casing, typos, mixed date and
price formats, duplicates), cleans it, assigns each line a category from a
keyword rule table, and builds a formula-driven summary workbook.

Usage:  python categorize.py
"""
import os, random, re
from datetime import date, timedelta
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
random.seed(21)

RULES = [  # (category, keywords) - first match wins
    ("Electronics", ["headphone", "earbud", "charger", "usb", "speaker", "power bank", "hdmi"]),
    ("Home & Kitchen", ["mug", "kettle", "pan", "knife", "towel", "candle", "blender"]),
    ("Fitness", ["yoga", "dumbbell", "resistance", "skipping", "water bottle", "foam roller"]),
    ("Clothing", ["t-shirt", "tshirt", "hoodie", "socks", "cap", "jacket", "leggings"]),
    ("Beauty", ["serum", "shampoo", "moisturi", "lip balm", "sunscreen"]),
    ("Pet Supplies", ["dog", "cat ", "pet ", "leash", "litter"]),
]
PRODUCTS = {
    "Electronics": [("Wireless Headphones", 59.99), ("USB-C Charger 30W", 18.5), ("Bluetooth Speaker Mini", 34.0),
                    ("Power Bank 10000mAh", 27.99), ("HDMI Cable 2m", 9.99), ("Earbuds Sport", 42.0)],
    "Home & Kitchen": [("Ceramic Mug Set", 22.0), ("Electric Kettle 1.7L", 31.5), ("Non-stick Frying Pan", 28.0),
                       ("Chef Knife 8in", 39.0), ("Scented Candle Vanilla", 12.5), ("Hand Blender", 45.0)],
    "Fitness": [("Yoga Mat 6mm", 24.99), ("Dumbbell Pair 5kg", 36.0), ("Resistance Band Set", 15.99),
                ("Steel Water Bottle 750ml", 17.0), ("Foam Roller", 21.0), ("Skipping Rope", 8.99)],
    "Clothing": [("Cotton T-Shirt", 14.99), ("Zip Hoodie", 44.0), ("Ankle Socks 5pk", 11.0),
                 ("Baseball Cap", 13.5), ("Rain Jacket", 62.0), ("Gym Leggings", 29.0)],
    "Beauty": [("Vitamin C Serum", 19.99), ("Argan Shampoo", 13.0), ("Daily Moisturiser SPF15", 16.5),
               ("Lip Balm Trio", 7.5), ("Sunscreen SPF50", 14.0)],
    "Pet Supplies": [("Dog Chew Toy", 9.5), ("Cat Scratching Post", 33.0), ("Pet Grooming Brush", 11.99),
                     ("Reflective Leash", 15.0), ("Clumping Litter 10L", 12.0)],
    None: [("Gift Card", 25.0), ("Mystery Box", 30.0), ("Assorted Stickers", 4.5)],
}
FIRST = ["Ava", "Liam", "Zoe", "Noah", "Mia", "Ethan", "Lily", "Sipho", "Emma", "Ravi", "Chloe", "Ben", "Thandi", "Leo", "Ruby"]
LAST = ["Carter", "Naidoo", "Hughes", "Mokoena", "Fischer", "Patel", "Brooks", "Santos", "Meyer", "Khumalo"]
UNIT_FIXES = {"Usb": "USB", "Hdmi": "HDMI", "Spf": "SPF", "Mah": "mAh", "Ml": "ml", "5Kg": "5kg",
              "5Pk": "5pk", "8In": "8in", "6Mm": "6mm", "2M": "2m", "Non-Stick": "Non-stick"}
REGIONS = ["North", "South", "East", "West"]


def messy(name):
    r = random.random()
    if r < .15: return name.upper()
    if r < .30: return name.lower()
    if r < .38: return "  " + name + " "
    if r < .44: return name.replace("T-Shirt", "Tshirt").replace("Moisturiser", "Moisturizer")
    return name


def make_raw(n=320):
    customers = [f"{random.choice(FIRST)} {random.choice(LAST)}" for _ in range(70)]
    rows, start = [], date(2026, 4, 1)
    for i in range(n):
        cat = random.choices(list(PRODUCTS), weights=[22, 20, 16, 18, 12, 9, 3])[0]
        name, price = random.choice(PRODUCTS[cat])
        d = start + timedelta(days=random.randint(0, 182))
        ds = d.strftime(random.choice(["%Y-%m-%d", "%d/%m/%Y", "%d %b %Y"]))
        ps = random.choice([f"{price:.2f}", f"${price:.2f}", f"USD {price:.2f}"])
        rows.append([f"ORD-{5000+i}", ds, random.choice(customers), messy(name), random.randint(1, 4), ps, random.choice(REGIONS)])
    rows += random.sample(rows, 9)          # duplicate order lines
    random.shuffle(rows)
    return rows


def parse_date(s):
    from datetime import datetime
    for f in ("%Y-%m-%d", "%d/%m/%Y", "%d %b %Y"):
        try: return datetime.strptime(s, f).date()
        except ValueError: pass
    return None


def categorize(name):
    low = name.lower() + " "
    for cat, kws in RULES:
        if any(k in low for k in kws):
            return cat
    return "Uncategorized"


def style_header(ws, fill="1F3A5F"):
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=fill)
        c.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"


def main():
    raw = make_raw()
    wb = Workbook()
    r = wb.active; r.title = "Raw Export"
    r.append(["order_id", "date", "customer", "item", "qty", "price", "region"])
    for row in raw: r.append(row)
    style_header(r, "7F7F7F")

    seen, clean = set(), []
    for oid, ds, cust, item, qty, price, region in raw:
        if oid in seen: continue
        seen.add(oid)
        item_c = re.sub(r"\s+", " ", item).strip().title().replace("Tshirt", "T-Shirt").replace("Moisturizer", "Moisturiser")
        for wrong, right in UNIT_FIXES.items():
            item_c = item_c.replace(wrong, right)
        clean.append([oid, parse_date(ds), cust, item_c, qty, float(re.sub(r"[^\d.]", "", price)), region, categorize(item_c)])
    clean.sort(key=lambda x: x[1])

    c = wb.create_sheet("Categorized")
    c.append(["Order ID", "Date", "Customer", "Item", "Qty", "Unit Price", "Region", "Category", "Line Total", "Month"])
    for i, row in enumerate(clean, 2):
        c.append(row + [f"=E{i}*F{i}", f'=TEXT(B{i},"yyyy-mm")'])
        c[f"B{i}"].number_format = "yyyy-mm-dd"
        c[f"F{i}"].number_format = c[f"I{i}"].number_format = "$#,##0.00"
        if row[7] == "Uncategorized":
            for cell in c[i]: cell.fill = PatternFill("solid", fgColor="FFF2CC")
    style_header(c)
    for col, w in zip("ABCDEFGHIJ", [12, 12, 18, 28, 6, 11, 9, 16, 12, 10]):
        c.column_dimensions[col].width = w
    c.auto_filter.ref = c.dimensions
    n = len(clean) + 1
    rng = lambda col: f"Categorized!${col}$2:${col}${n}"

    s = wb.create_sheet("Summary")
    s["A1"] = "Purchase Summary by Category"; s["A1"].font = Font(bold=True, size=14)
    s["A2"] = "Sample project - fictional data. All figures are live formulas over the Categorized sheet."
    s["A2"].font = Font(italic=True, color="808080")
    s.append([]); s.append(["Category", "Orders", "Units", "Revenue", "Avg Order Value", "% of Revenue"])
    for cell in s[4]:
        cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="1F3A5F")
    cats = [x[0] for x in RULES] + ["Uncategorized"]
    for i, cat in enumerate(cats, 5):
        s.append([cat, f"=COUNTIF({rng('H')},A{i})", f"=SUMIF({rng('H')},A{i},{rng('E')})",
                  f"=SUMIF({rng('H')},A{i},{rng('I')})", f"=IF(B{i}=0,0,D{i}/B{i})", f"=D{i}/$D${5+len(cats)}"])
    t = 5 + len(cats)
    s.append(["Total", f"=SUM(B5:B{t-1})", f"=SUM(C5:C{t-1})", f"=SUM(D5:D{t-1})", f"=D{t}/B{t}", f"=SUM(F5:F{t-1})"])
    for cell in s[t]: cell.font = Font(bold=True)
    for i in range(5, t + 1):
        s[f"D{i}"].number_format = s[f"E{i}"].number_format = "$#,##0.00"; s[f"F{i}"].number_format = "0.0%"

    s[f"A{t+2}"] = "Revenue by Region"; s[f"A{t+2}"].font = Font(bold=True, size=12)
    hr = t + 3
    for j, h in enumerate(["Category"] + REGIONS, 1):
        cell = s.cell(row=hr, column=j, value=h); cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="1F3A5F")
    for i, cat in enumerate(cats, hr + 1):
        s.cell(row=i, column=1, value=cat)
        for j, reg in enumerate(REGIONS, 2):
            col = get_column_letter(j)
            s.cell(row=i, column=j, value=f"=SUMIFS({rng('I')},{rng('H')},$A{i},{rng('G')},{col}${hr})").number_format = "$#,##0"
    for col, w in zip("ABCDEF", [18, 10, 10, 14, 16, 14]): s.column_dimensions[col].width = w

    ch = BarChart(); ch.type = "bar"; ch.title = "Revenue by Category"; ch.legend = None
    ch.add_data(Reference(s, min_col=4, min_row=4, max_row=t - 1), titles_from_data=True)
    ch.set_categories(Reference(s, min_col=1, min_row=5, max_row=t - 1))
    ch.height, ch.width = 8, 16
    s.add_chart(ch, "H4")

    ru = wb.create_sheet("Category Rules")
    ru.append(["Category", "Keywords (first match wins)"])
    for cat, kws in RULES: ru.append([cat, ", ".join(kws)])
    ru.append(["Uncategorized", "anything that matches no rule - highlighted yellow for manual review"])
    style_header(ru); ru.column_dimensions["A"].width = 18; ru.column_dimensions["B"].width = 80

    wb.move_sheet("Summary", offset=-2)
    out = os.path.join(HERE, "purchase_data_categorized.xlsx")
    wb.save(out)
    unc = sum(1 for x in clean if x[7] == "Uncategorized")
    rev = sum(x[4] * x[5] for x in clean)
    print(f"raw rows {len(raw)} -> clean {len(clean)} (removed {len(raw)-len(clean)} duplicates), uncategorized {unc}, revenue {rev:.2f}")


if __name__ == "__main__":
    main()
