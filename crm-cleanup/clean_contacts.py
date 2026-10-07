"""
CRM contact list clean-up and de-duplication (sample project, fictional data)

Takes a messy contact export and returns a clean, import-ready list:
names and companies in consistent casing, phones in one format, emails
validated, duplicates merged (keeping the most complete record), and a
log of every change so nothing is altered silently.

Usage:  python clean_contacts.py
"""
import os, random, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

HERE = os.path.dirname(os.path.abspath(__file__))
random.seed(5)
EMAIL_RE = re.compile(r"^[\w.+-]+@[\w-]+(\.[\w-]+)+$")
SUFFIXES = {"ltd": "Ltd", "llc": "LLC", "inc": "Inc", "co": "Co", "pty": "Pty"}


def make_raw():
    first = ["olivia", "james", "aisha", "daniel", "mei", "carlos", "hannah", "tunde", "sara", "luke", "nina", "omar", "grace", "ivan", "zara", "ben"]
    last = ["bennett", "o'connor", "khan", "mcdonald", "chen", "ramirez", "schmidt", "adeyemi", "lindqvist", "foster", "petrova", "hassan"]
    comps = ["brightway logistics ltd", "harper & lane", "nova dental", "kingfisher media inc", "summit accounting", "oakridge property co",
             "bluebird travel", "ironwood construction llc", "lumen health", "pioneer print"]
    cities = ["london", "manchester", "leeds", "bristol", "glasgow", "cardiff"]
    rows = []
    for i in range(140):
        f, l, c = random.choice(first), random.choice(last), random.choice(comps)
        phone = "7" + "".join(random.choice("0123456789") for _ in range(9))
        email = f"{f}.{l.replace(chr(39), '')}@{re.sub(r'[^a-z]', '', c)[:12]}.example"
        rows.append([f, l, email, phone, c, random.choice(cities), random.choice(["Lead", "Customer", "Customer", "Prospect"])])
    out = []
    for f, l, email, phone, c, city, status in rows:
        f, l, c, city = (random.choice([str.upper, str.lower, str.title, str.title])(x) for x in (f, l, c, city))
        phone = random.choice(["+44 {}", "0{}", "44{}", "(0){}", "+44{}"]).format(phone)
        if random.random() < .3: phone = phone[:-6] + " " + phone[-6:-3] + "-" + phone[-3:]
        if random.random() < .15: email = email.upper()
        if random.random() < .08: email = " " + email + " "
        r = random.random()
        if r < .04: email = email.replace("@", "(at)")
        elif r < .07: email = ""
        if random.random() < .05: phone = phone[:6]
        if random.random() < .1: city = ""
        out.append([f"{f} {l}" if random.random() < .9 else f"{l}, {f}", email, phone, c, city, status])
    for row in random.sample(out[:], 22):          # duplicates, some with gaps or different casing
        d = row[:]
        d[0] = d[0].upper(); d[4] = "" if random.random() < .5 else d[4]
        out.append(d)
    random.shuffle(out)
    return out


def clean_name(n):
    n = re.sub(r"\s+", " ", n).strip()
    if "," in n:
        last, first = [p.strip() for p in n.split(",", 1)]; n = f"{first} {last}"
    n = n.title()
    n = re.sub(r"\bMc(\w)", lambda m: "Mc" + m.group(1).upper(), n)
    return re.sub(r"O'(\w)", lambda m: "O'" + m.group(1).upper(), n)


def clean_company(c):
    words = re.sub(r"\s+", " ", c).strip().title().split(" ")
    return " ".join(SUFFIXES.get(w.lower(), w) for w in words)


def clean_phone(p):
    d = re.sub(r"\D", "", p)
    if d.startswith("44"): d = d[2:]
    d = d.lstrip("0")
    return (f"+44 {d[:4]} {d[4:]}", True) if len(d) == 10 else (p.strip(), False)


def main():
    raw = make_raw()
    cleaned, log = [], {"Names re-cased or reordered": 0, "Companies standardised": 0, "Phones reformatted": 0,
                        "Emails trimmed / lower-cased": 0, "Invalid or missing email flagged": 0, "Invalid phone flagged": 0}
    for name, email, phone, comp, city, status in raw:
        n, c, e = clean_name(name), clean_company(comp), email.strip().lower()
        p, p_ok = clean_phone(phone)
        e_ok = bool(EMAIL_RE.match(e))
        log["Names re-cased or reordered"] += n != name
        log["Companies standardised"] += c != comp
        log["Phones reformatted"] += p_ok and p != phone
        log["Emails trimmed / lower-cased"] += e_ok and e != email
        log["Invalid or missing email flagged"] += not e_ok
        log["Invalid phone flagged"] += not p_ok
        issues = [] if e_ok else ["Email missing" if not e else "Email invalid"]
        if not p_ok: issues.append("Phone invalid")
        if not city: issues.append("City missing")
        cleaned.append({"Name": n, "Email": e, "Phone": p, "Company": c, "City": city.title(), "Status": status, "Review": "; ".join(issues) or "OK",
                        "_key": e if e_ok else (p if p_ok else None)})

    # de-duplicate on email (or phone when there is no usable email); keep the most complete record
    best, removed = {}, []
    filled = lambda r: sum(bool(r[k]) for k in ("Email", "Phone", "City")) + (r["Review"] == "OK")
    for r in cleaned:
        k = r["_key"] or id(r)
        if k in best:
            keep, drop = (r, best[k]) if filled(r) > filled(best[k]) else (best[k], r)
            for f in ("City", "Phone", "Email"):
                if not keep[f] and drop[f]: keep[f] = drop[f]
            best[k] = keep; removed.append({**drop, "Kept as": keep["Name"]})
        else:
            best[k] = r
    final = sorted(best.values(), key=lambda r: (r["Company"], r["Name"]))
    for r in final:      # merged gaps may have fixed a flag
        if r["City"] and "City missing" in r["Review"]:
            r["Review"] = "; ".join(x for x in r["Review"].split("; ") if x != "City missing") or "OK"

    wb = Workbook()
    head = lambda ws, color: [setattr(c, "font", Font(bold=True, color="FFFFFF")) or setattr(c, "fill", PatternFill("solid", fgColor=color)) for c in ws[1]]
    cols = ["Name", "Email", "Phone", "Company", "City", "Status", "Review"]

    s = wb.active; s.title = "Summary"
    s["A1"] = "Contact List Clean-up"; s["A1"].font = Font(bold=True, size=14)
    s["A2"] = "Sample project - fictional data"; s["A2"].font = Font(italic=True, color="808080")
    n = len(final) + 1
    stats = [("Rows received", len(raw)), ("Duplicates merged", len(removed)), ("Clean contacts delivered", f"=COUNTA('Clean Contacts'!A2:A{n})"),
             ("Ready to import (no flags)", f"=COUNTIF('Clean Contacts'!G2:G{n},\"OK\")"), ("Need a manual look", "=B6-B7")]
    for i, (k, v) in enumerate(stats, 4): s[f"A{i}"], s[f"B{i}"] = k, v
    s["A10"] = "Changes made"; s["A10"].font = Font(bold=True)
    for i, (k, v) in enumerate(log.items(), 11): s[f"A{i}"], s[f"B{i}"] = k, v
    s["A18"] = "Contacts by status"; s["A18"].font = Font(bold=True)
    for i, st in enumerate(["Customer", "Lead", "Prospect"], 19):
        s[f"A{i}"], s[f"B{i}"] = st, f"=COUNTIF('Clean Contacts'!F2:F{n},A{i})"
    s.column_dimensions["A"].width = 34; s.column_dimensions["B"].width = 12

    c = wb.create_sheet("Clean Contacts"); c.append(cols)
    for r in final:
        c.append([r[k] for k in cols])
        if r["Review"] != "OK":
            for cell in c[c.max_row]: cell.fill = PatternFill("solid", fgColor="FFF2CC")
    head(c, "1F3A5F"); c.freeze_panes = "A2"; c.auto_filter.ref = c.dimensions
    d = wb.create_sheet("Duplicates Removed"); d.append(cols[:-1] + ["Merged into"])
    for r in removed: d.append([r[k] for k in cols[:-1]] + [r["Kept as"]])
    head(d, "7A3E3E"); d.freeze_panes = "A2"
    o = wb.create_sheet("Original Export"); o.append(["name", "email", "phone", "company", "city", "status"])
    for r in raw: o.append(r)
    head(o, "7F7F7F"); o.freeze_panes = "A2"
    for ws in (c, d, o):
        for col, w in zip("ABCDEFG", [22, 40, 20, 28, 14, 12, 28]): ws.column_dimensions[col].width = w

    wb.save(os.path.join(HERE, "contacts_cleaned.xlsx"))
    flagged = sum(r["Review"] != "OK" for r in final)
    print(f"raw {len(raw)} -> clean {len(final)}, duplicates merged {len(removed)}, flagged {flagged}")
    print(log)
    for r in final[:4]: print({k: r[k] for k in cols})


if __name__ == "__main__":
    main()
