"""
Ecommerce product listings (sample project, fictional store "Hearth & Hollow")

Builds two files from one set of written listings:
  product_listings.xlsx      - review sheet with titles, descriptions, bullets, SEO fields and length checks
  shopify_import.csv         - the same products in Shopify's product CSV layout, ready to import

Usage:  python build_listings.py
"""
import csv, os, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

HERE = os.path.dirname(os.path.abspath(__file__))
VENDOR = "Hearth & Hollow"

# name, type, price, compare_at, variants, description, bullets, tags
P = [
 ("Stoneware Pour-Over Set", "Coffee & Tea", 46, None, ["Oat", "Charcoal"],
  "Slow coffee without the fuss. The dripper sits straight on the matching mug, so there is one less thing to wash and nothing to balance.",
  ["Hand-glazed stoneware that holds heat while you pour", "Dripper takes standard size 2 paper filters", "Mug holds 350 ml", "Dishwasher safe"],
  "pour over, coffee dripper, stoneware mug, coffee gift"),
 ("Linen Kitchen Apron", "Kitchen Textiles", 38, 48, ["Clay", "Sage", "Ink"],
  "A proper apron that gets softer every wash. Cross-back straps take the weight off your neck, and the front pocket is big enough for a phone and a tea towel.",
  ["100% European flax linen", "Cross-back straps, one size fits most", "Deep front pocket", "Machine wash cold, line dry"],
  "linen apron, cross back apron, kitchen apron, baking gift"),
 ("Acacia Serving Board", "Serveware", 34, None, ["Medium", "Large"],
  "One board for bread, cheese and everything you put out when people come over. Each piece of acacia has its own grain, so no two look the same.",
  ["Solid acacia wood, food-safe oil finish", "Medium 35 x 20 cm, Large 45 x 25 cm", "Handle with hanging loop", "Hand wash only"],
  "serving board, cheese board, acacia wood, hosting"),
 ("Beeswax Food Wraps, Set of 3", "Food Storage", 18, None, ["Meadow print", "Plain"],
  "Wrap half a lemon, cover a bowl, pack a sandwich. Warm the wrap in your hands and it holds its shape, then rinse and use it again.",
  ["Organic cotton coated in beeswax and jojoba oil", "Three sizes: small, medium, large", "Lasts around a year with regular use", "Rinse in cool water"],
  "beeswax wraps, reusable food wrap, plastic free, zero waste kitchen"),
 ("Cast Iron Trivet", "Kitchen Tools", 24, None, ["Matte black"],
  "Heavy enough to stay put under a full casserole dish, good-looking enough to leave on the table. Rubber feet keep it from scratching the surface underneath.",
  ["Solid cast iron, 18 cm across", "Heat safe to oven temperatures", "Four non-slip rubber feet", "Wipe clean"],
  "trivet, cast iron, pot stand, table protector"),
 ("Soy Candle, Fig & Cedar", "Home Fragrance", 22, 28, ["200 g"],
  "Green fig, warm cedar and a little black pepper. A scent that fills the room without taking it over.",
  ["Natural soy wax, cotton wick", "About 40 hours burn time", "Poured into a reusable glass jar", "Trim wick to 5 mm before lighting"],
  "soy candle, fig candle, scented candle, housewarming gift"),
 ("Waffle Cotton Hand Towels, Pair", "Bath Linen", 26, None, ["Stone", "Rust", "White"],
  "Light, quick to dry and far more absorbent than they look. The waffle weave gets fluffier after the first wash.",
  ["100% organic cotton, 50 x 80 cm", "Hanging loop on each towel", "Dries faster than terry", "Machine wash warm"],
  "hand towels, waffle towel, organic cotton, bathroom"),
 ("Glass Storage Jars, Set of 4", "Food Storage", 32, None, ["500 ml", "1 litre"],
  "See what you have at a glance. The bamboo lid seals with a silicone ring, so flour stays dry and coffee stays fresh.",
  ["Borosilicate glass with bamboo lids", "Airtight silicone seal", "Stackable", "Jars are dishwasher safe, hand wash lids"],
  "storage jars, pantry organisation, glass jars, bamboo lid"),
]
SITE = "hearthandhollow.example"
slug = lambda s: re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")


def seo_title(name, ptype):
    t = f"{name} | {VENDOR}"
    return t if len(t) <= 60 else name


def meta(desc):
    first = desc.split(". ")[0].rstrip(".") + "."
    m = f"{first} Free shipping over $60."
    return m if len(m) <= 155 else first


def main():
    wb = Workbook(); ws = wb.active; ws.title = "Listings"
    cols = ["Product", "Type", "Price", "Compare at", "Variants", "Description", "Bullet points", "SEO title", "Title length",
            "Meta description", "Meta length", "URL handle", "Tags", "Check"]
    ws.append(cols)
    for i, (name, ptype, price, was, variants, desc, bullets, tags) in enumerate(P, 2):
        ws.append([name, ptype, price, was, ", ".join(variants), desc, "\n".join("- " + b for b in bullets), seo_title(name, ptype),
                   f"=LEN(H{i})", meta(desc), f"=LEN(J{i})", slug(name), tags,
                   f'=IF(AND(I{i}<=60,K{i}<=155),"OK","Too long")'])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="3D405B")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col, w in zip("ABCDEFGHIJKLMN", [28, 16, 9, 11, 20, 52, 46, 38, 9, 52, 9, 30, 38, 10]):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
        row[2].number_format = row[3].number_format = "$#,##0.00"
    ws.freeze_panes = "B2"; ws.row_dimensions[1].height = 30

    g = wb.create_sheet("Style Guide")
    for line in ["Listing style guide (sample project - fictional store)", "",
                 "Product title: what it is first, then material or set size. No ALL CAPS, no keyword stuffing.",
                 "Description: 2 to 3 sentences. Start with how it is used, not with the brand.",
                 "Bullets: four per product. Material, size, one practical benefit, care.",
                 "SEO title: 60 characters or fewer. Meta description: 155 or fewer. The Check column flags anything over.",
                 "Claims: only what the supplier spec sheet supports. No 'best', 'premium quality' or invented numbers.",
                 "Tags: four per product, lower case, the phrases shoppers actually type."]:
        g.append([line])
    g["A1"].font = Font(bold=True, size=13); g.column_dimensions["A"].width = 110
    wb.save(os.path.join(HERE, "product_listings.xlsx"))

    head = ["Handle", "Title", "Body (HTML)", "Vendor", "Type", "Tags", "Published", "Option1 Name", "Option1 Value", "Variant SKU",
            "Variant Inventory Policy", "Variant Fulfillment Service", "Variant Price", "Variant Compare At Price",
            "Variant Requires Shipping", "Variant Taxable", "SEO Title", "SEO Description", "Status"]
    rows = 0
    with open(os.path.join(HERE, "shopify_import.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f); w.writerow(head)
        for n, (name, ptype, price, was, variants, desc, bullets, tags) in enumerate(P, 1):
            body = f"<p>{desc}</p><ul>" + "".join(f"<li>{b}</li>" for b in bullets) + "</ul>"
            opt = "Size" if any(ch.isdigit() for ch in variants[0]) or variants[0] in ("Medium", "Large") else ("Style" if "print" in variants[0] else "Colour")
            for vi, v in enumerate(variants):
                sku = f"HH-{n:02d}-{re.sub(r'[^A-Z0-9]', '', v.upper())[:4]}"
                first = vi == 0   # Shopify wants product-level fields on the first variant row only
                w.writerow([slug(name), name if first else "", body if first else "", VENDOR if first else "", ptype if first else "",
                            tags if first else "", "TRUE" if first else "", opt if first else "", v, sku, "deny", "manual",
                            f"{price:.2f}", f"{was:.2f}" if was else "", "TRUE", "TRUE",
                            seo_title(name, ptype) if first else "", meta(desc) if first else "", "draft" if first else ""])
                rows += 1
    print(f"{len(P)} products, {rows} variant rows")
    for p in P: print(len(seo_title(p[0], p[1])), len(meta(p[5])), seo_title(p[0], p[1]))


if __name__ == "__main__":
    main()
