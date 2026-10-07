"""
Bulk Excel -> branded PDF statements (sample project, fictional data)

Reads one workbook of invoice lines and produces one formatted PDF
statement per client, plus an index sheet listing what was generated.

Usage:  python excel_to_pdfs.py
"""
import os, random
from collections import defaultdict
from datetime import date, timedelta
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "invoice_lines.xlsx")
OUT = os.path.join(HERE, "output_pdfs")
BRAND, ACCENT = colors.HexColor("#0F4C5C"), colors.HexColor("#E36414")


def make_source():
    random.seed(11)
    clients = [("Brightside Dental", "12 Oak Avenue"), ("Corner Bakehouse", "88 Mill Street"),
               ("Lumen Yoga Studio", "5 Harbour Road"), ("Paws & Co Veterinary", "41 Garden Lane"),
               ("Tidewater Law", "200 Market Square"), ("Greenleaf Nursery", "7 Orchard Way"),
               ("Atlas Auto Repair", "63 Station Road"), ("Pixel & Pine Studio", "19 Kings Parade")]
    services = [("Monthly bookkeeping", 320), ("Payroll processing", 180), ("Data entry (per 100 records)", 45),
                ("Inbox management", 150), ("Customer support cover", 260), ("Report preparation", 95),
                ("CRM clean-up", 120), ("Website content update", 85)]
    wb = Workbook(); ws = wb.active; ws.title = "Invoice Lines"
    ws.append(["Client", "Address", "Invoice No", "Date", "Service", "Qty", "Unit Price", "Status"])
    inv = 3001
    for name, addr in clients:
        for _ in range(random.randint(3, 6)):
            svc, price = random.choice(services)
            d = date(2026, 7, 1) + timedelta(days=random.randint(0, 85))
            ws.append([name, addr, f"INV-{inv}", d, svc, random.randint(1, 4), price, random.choice(["Paid", "Paid", "Unpaid"])])
            inv += 1
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="0F4C5C")
    for col, w in zip("ABCDEFGH", [24, 20, 12, 12, 32, 6, 11, 10]): ws.column_dimensions[col].width = w
    for r in ws.iter_rows(min_row=2):
        r[3].number_format = "yyyy-mm-dd"; r[6].number_format = "$#,##0.00"
    wb.save(SRC)


def build_pdf(client, addr, lines, path):
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=18*mm, rightMargin=18*mm, topMargin=16*mm, bottomMargin=18*mm)
    h1 = ParagraphStyle("h1", fontName="Helvetica-Bold", fontSize=20, textColor=BRAND, leading=24)
    small = ParagraphStyle("s", fontName="Helvetica", fontSize=9, textColor=colors.grey, leading=12)
    body = ParagraphStyle("b", fontName="Helvetica", fontSize=10.5, leading=14)
    total = sum(l[3] * l[4] for l in lines)
    unpaid = sum(l[3] * l[4] for l in lines if l[5] == "Unpaid")
    story = [Paragraph("Northwind Business Services", h1),
             Paragraph("Fictional company - sample statement for demonstration", small), Spacer(1, 8*mm),
             Paragraph(f"<b>Statement for:</b> {client.replace('&', '&amp;')}<br/>{addr}", body),
             Paragraph("<b>Period:</b> 1 July 2026 - 30 September 2026", body), Spacer(1, 6*mm)]
    data = [["Invoice", "Date", "Service", "Qty", "Unit", "Amount", "Status"]]
    for no, d, svc, q, p, st in sorted(lines, key=lambda x: x[1]):
        data.append([no, d.strftime("%d %b %Y"), svc, q, f"${p:,.2f}", f"${q*p:,.2f}", st])
    t = Table(data, colWidths=[22*mm, 24*mm, 58*mm, 10*mm, 20*mm, 22*mm, 18*mm], repeatRows=1)
    st = [("BACKGROUND", (0, 0), (-1, 0), BRAND), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
          ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 9),
          ("ALIGN", (3, 0), (5, -1), "RIGHT"), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F6F7")]),
          ("LINEBELOW", (0, 0), (-1, -1), 0.3, colors.HexColor("#D5DCDE")), ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]
    for i, row in enumerate(data[1:], 1):
        if row[6] == "Unpaid":
            st += [("TEXTCOLOR", (6, i), (6, i), ACCENT), ("FONTNAME", (6, i), (6, i), "Helvetica-Bold")]
    t.setStyle(TableStyle(st))
    tot = Table([["Total billed", f"${total:,.2f}"], ["Paid", f"${total-unpaid:,.2f}"], ["Balance due", f"${unpaid:,.2f}"]],
                colWidths=[40*mm, 30*mm], hAlign="RIGHT")
    tot.setStyle(TableStyle([("ALIGN", (1, 0), (1, -1), "RIGHT"), ("FONTSIZE", (0, 0), (-1, -1), 10),
                             ("FONTNAME", (0, 2), (-1, 2), "Helvetica-Bold"), ("LINEABOVE", (0, 2), (-1, 2), 1, BRAND),
                             ("TEXTCOLOR", (0, 2), (-1, 2), ACCENT if unpaid else BRAND)]))
    story += [t, Spacer(1, 6*mm), tot, Spacer(1, 10*mm),
              Paragraph("Thank you for your business. Please quote the invoice number with any payment.", small)]
    doc.build(story)
    return total, unpaid


def main():
    make_source()
    os.makedirs(OUT, exist_ok=True)
    ws = load_workbook(SRC)["Invoice Lines"]
    groups, addrs = defaultdict(list), {}
    for client, addr, no, d, svc, q, p, st in ws.iter_rows(min_row=2, values_only=True):
        groups[client].append((no, d, svc, q, p, st)); addrs[client] = addr
    wb = load_workbook(SRC)
    idx = wb.create_sheet("PDF Index")
    idx.append(["Client", "PDF File", "Invoices", "Total Billed", "Balance Due"])
    for c in idx[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="0F4C5C")
    for client, lines in groups.items():
        fn = "statement_" + "".join(ch if ch.isalnum() else "_" for ch in client).strip("_").replace("___", "_").lower() + ".pdf"
        total, unpaid = build_pdf(client, addrs[client], lines, os.path.join(OUT, fn))
        idx.append([client, fn, len(lines), total, unpaid])
        print(f"{fn:48s} {len(lines)} lines  total {total:>9,.2f}  due {unpaid:>9,.2f}")
    for r in idx.iter_rows(min_row=2):
        r[3].number_format = r[4].number_format = "$#,##0.00"
    for col, w in zip("ABCDE", [26, 46, 10, 14, 14]): idx.column_dimensions[col].width = w
    wb.save(SRC)


if __name__ == "__main__":
    main()
