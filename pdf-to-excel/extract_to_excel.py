"""
PDF form -> Excel extraction (sample project, fictional data)

Reads every PDF in input_pdfs/, pulls the labelled fields, cleans them,
flags anything that needs a human look, and writes one formatted workbook.

Usage:  python extract_to_excel.py
"""
import os, re, glob
import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
FIELDS = ["Reference No", "Date Submitted", "Company Name", "Contact Person", "Email", "Phone",
          "Category", "Payment Terms", "Annual Turnover (USD)", "Tax Registered"]
EMAIL_RE = re.compile(r"^[\w.+-]+@[\w-]+(\.[\w-]+)+$")


def read_form(path):
    with pdfplumber.open(path) as pdf:
        text = "\n".join(p.extract_text() or "" for p in pdf.pages)
    row = {}
    for f in FIELDS:
        m = re.search(rf"^{re.escape(f)}:[ \t]*(.*)$", text, re.M)
        row[f] = m.group(1).strip() if m else ""
    return row


def clean(row):
    issues = []
    row["Company Name"] = row["Company Name"].title().replace(" Co", " Co") if row["Company Name"].isupper() else row["Company Name"]
    row["Email"] = row["Email"].replace(" ", "").lower()
    if not row["Email"]:
        issues.append("Email missing")
    elif not EMAIL_RE.match(row["Email"]):
        issues.append("Email format")
    digits = re.sub(r"\D", "", row["Phone"])
    if len(digits) == 11:
        row["Phone"] = f"+{digits[:2]} {digits[2:4]} {digits[4:7]} {digits[7:]}"
    else:
        issues.append("Phone format")
    try:
        row["Annual Turnover (USD)"] = int(row["Annual Turnover (USD)"].replace(",", ""))
    except ValueError:
        issues.append("Turnover not numeric")
    d = row["Date Submitted"].split("/")
    if len(d) == 3:
        row["Date Submitted"] = f"{d[2]}-{d[1]}-{d[0]}"
    for f in FIELDS:
        if row[f] == "" and f != "Email":
            issues.append(f"{f} missing")
    row["Review Flag"] = "; ".join(issues) if issues else "OK"
    return row


def main():
    files = sorted(glob.glob(os.path.join(HERE, "input_pdfs", "*.pdf")))
    rows = []
    for f in files:
        r = clean(read_form(f))
        r["Source File"] = os.path.basename(f)
        rows.append(r)

    wb = Workbook()
    ws = wb.active
    ws.title = "Extracted Data"
    cols = FIELDS + ["Review Flag", "Source File"]
    head_fill = PatternFill("solid", fgColor="1F3A5F")
    thin = Side(style="thin", color="D0D0D0")
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = head_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r in rows:
        ws.append([r[c] for c in cols])
    flag_col = cols.index("Review Flag") + 1
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.border = Border(bottom=thin)
        if row[flag_col - 1].value != "OK":
            for c in row:
                c.fill = PatternFill("solid", fgColor="FFF2CC")
        row[cols.index("Annual Turnover (USD)")].number_format = "#,##0"
    for i, c in enumerate(cols, 1):
        width = max(len(str(c)), *(len(str(r[c])) for r in rows)) + 3
        ws.column_dimensions[get_column_letter(i)].width = min(width, 34)
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    n = len(rows) + 1
    s = wb.create_sheet("Summary")
    s["A1"] = "Extraction Summary"; s["A1"].font = Font(bold=True, size=14)
    s["A2"] = "Sample project - fictional data"; s["A2"].font = Font(italic=True, color="808080")
    tl = get_column_letter(cols.index("Annual Turnover (USD)") + 1)
    fl = get_column_letter(flag_col)
    cl = get_column_letter(cols.index("Category") + 1)
    items = [("Forms processed", f"=COUNTA('Extracted Data'!A2:A{n})"),
             ("Clean records", f"=COUNTIF('Extracted Data'!{fl}2:{fl}{n},\"OK\")"),
             ("Records needing review", "=B4-B5"),
             ("Total annual turnover (USD)", f"=SUM('Extracted Data'!{tl}2:{tl}{n})"),
             ("Average turnover (USD)", f"=AVERAGE('Extracted Data'!{tl}2:{tl}{n})")]
    for i, (k, v) in enumerate(items, 4):
        s[f"A{i}"] = k; s[f"B{i}"] = v; s[f"B{i}"].number_format = "#,##0"
    s["A11"] = "Vendors by category"; s["A11"].font = Font(bold=True)
    s["A12"], s["B12"], s["C12"] = "Category", "Vendors", "Turnover (USD)"
    for c in s[12]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = head_fill
    for i, cat in enumerate(sorted({r["Category"] for r in rows}), 13):
        s[f"A{i}"] = cat
        s[f"B{i}"] = f"=COUNTIF('Extracted Data'!{cl}2:{cl}{n},A{i})"
        s[f"C{i}"] = f"=SUMIF('Extracted Data'!{cl}2:{cl}{n},A{i},'Extracted Data'!{tl}2:{tl}{n})"
        s[f"C{i}"].number_format = "#,##0"
    s.column_dimensions["A"].width = 32; s.column_dimensions["B"].width = 16; s.column_dimensions["C"].width = 18

    out = os.path.join(HERE, "vendor_forms_extracted.xlsx")
    wb.save(out)
    flagged = [r for r in rows if r["Review Flag"] != "OK"]
    print(f"{len(rows)} forms -> {out}")
    for r in flagged:
        print("  review:", r["Source File"], "-", r["Review Flag"])


if __name__ == "__main__":
    main()
